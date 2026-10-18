"""Tests for reading the close inputs from an .xlsx workbook, and writing one."""

from __future__ import annotations

import shutil
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from intercompany_close_agent.cli import app
from intercompany_close_agent.ingest.loaders import DataLoadError, _cell, input_files, load_dataset
from intercompany_close_agent.ingest.workbook import write_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data" / "synthetic"
TEMPLATE = ROOT / "examples" / "close_template.xlsx"


def _edited(tmp_path: Path, edit) -> Path:
    """A copy of the template with one change made in Excel's terms."""
    path = tmp_path / "close.xlsx"
    shutil.copyfile(TEMPLATE, path)
    workbook = load_workbook(path)
    edit(workbook)
    workbook.save(path)
    return path


def test_the_committed_template_is_the_sample_data():
    """Same entities, balances and legs, down to how each amount is written:
    a whole amount comes back from Excel as an int and is padded like the CSV."""
    from_csv = load_dataset(DATA_DIR)
    from_workbook = load_dataset(TEMPLATE)

    assert from_workbook == from_csv
    assert [str(t.amount) for t in from_workbook[2]] == [str(t.amount) for t in from_csv[2]]
    assert [str(line.balance) for line in from_workbook[1]] == [str(line.balance) for line in from_csv[1]]


def test_the_template_is_what_write_workbook_makes_from_the_sample(tmp_path: Path):
    written = write_workbook(DATA_DIR, tmp_path / "close.xlsx")

    assert load_dataset(written) == load_dataset(TEMPLATE)


def test_a_run_from_the_workbook_prints_what_a_run_from_the_csvs_does(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_MODEL", raising=False)
    runner = CliRunner()

    from_csv = runner.invoke(app, ["run", "--data-dir", str(DATA_DIR), "--prepared-by", "Ana Ruiz"])
    from_workbook = runner.invoke(app, ["run", "--data", str(TEMPLATE), "--prepared-by", "Ana Ruiz"])

    assert from_csv.exit_code == from_workbook.exit_code == 0
    assert from_workbook.output == from_csv.output


def test_a_bad_cell_is_named_by_sheet_and_row(tmp_path: Path):
    def bad_amount(workbook):
        workbook["ic_ledger"]["D3"] = "abc"

    with pytest.raises(DataLoadError, match=r"close\.xlsx, sheet ic_ledger, row 3: could not parse amount: 'abc'"):
        load_dataset(_edited(tmp_path, bad_amount))


def test_a_missing_sheet_is_named(tmp_path: Path):
    def drop_ledger(workbook):
        workbook.remove(workbook["ic_ledger"])

    with pytest.raises(DataLoadError, match="close.xlsx has no sheet named ic_ledger"):
        load_dataset(_edited(tmp_path, drop_ledger))


def test_a_missing_column_is_named_with_its_sheet(tmp_path: Path):
    def drop_currency(workbook):
        workbook["entities"].delete_cols(4)

    with pytest.raises(DataLoadError, match="sheet entities, is missing column"):
        load_dataset(_edited(tmp_path, drop_currency))


def test_an_empty_ledger_sheet_is_nothing_to_close(tmp_path: Path):
    def empty_ledger(workbook):
        sheet = workbook["ic_ledger"]
        sheet.delete_rows(2, sheet.max_row)

    with pytest.raises(DataLoadError, match="sheet ic_ledger, has no transactions, so there is nothing to close"):
        load_dataset(_edited(tmp_path, empty_ledger))


def test_blank_rows_are_skipped_and_typed_values_read_as_the_csv_would(tmp_path: Path):
    """What a person does in Excel: leave a blank row, type an amount without
    decimals, type a date as text."""

    def hand_edits(workbook):
        sheet = workbook["ic_ledger"]
        sheet.insert_rows(3)
        sheet["D2"] = 83357
        sheet["F2"] = "2025-09-22"

    entities, trial_balance, transactions = load_dataset(_edited(tmp_path, hand_edits))

    assert len(transactions) == len(load_dataset(DATA_DIR)[2])
    assert str(transactions[0].amount) == "83357.00"
    assert transactions[0].transaction_date == date(2025, 9, 22)


def test_a_file_that_is_not_a_workbook_says_so(tmp_path: Path):
    fake = tmp_path / "close.xlsx"
    fake.write_text("entity_code,amount\n", encoding="utf-8")

    with pytest.raises(DataLoadError, match="isn't a readable .xlsx workbook"):
        load_dataset(fake)


@pytest.mark.parametrize(
    ("value", "column", "text"),
    [
        (None, "amount", ""),
        (True, "interfaced", "True"),
        (123964, "amount", "123964.00"),
        (1000.0, "balance", "1000.00"),
        (83357.72, "amount", "83357.72"),
        (0.005, "amount", "0.005"),
        (1004, "reference", "1004"),
        (datetime(2025, 9, 30, 0, 0), "transaction_date", "2025-09-30"),
        ("RECHARGE-2025-1000", "reference", "RECHARGE-2025-1000"),
    ],
)
def test_a_cell_reads_as_the_text_its_csv_field_would_have(value, column: str, text: str):
    assert _cell(value, column) == text


def test_a_workbook_is_one_input_file_and_a_folder_is_three():
    assert input_files(TEMPLATE) == [TEMPLATE]
    assert [p.name for p in input_files(DATA_DIR)] == ["entities.csv", "trial_balance.csv", "ic_ledger.csv"]


def test_write_workbook_command_writes_a_workbook_run_can_read(tmp_path: Path):
    out = tmp_path / "inputs.xlsx"

    result = CliRunner().invoke(app, ["write-workbook", "--data-dir", str(DATA_DIR), "--out", str(out)])

    assert result.exit_code == 0, result.output
    assert load_dataset(out) == load_dataset(DATA_DIR)
