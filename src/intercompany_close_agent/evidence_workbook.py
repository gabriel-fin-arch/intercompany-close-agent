"""The evidence pack as one workbook: close.xlsx, for a reviewer who works in Excel.

A first sheet with what run.json says — who prepared and who reviewed, when,
the policy and the code version — then one sheet per CSV the pack writes:
tie-out, exceptions, decisions, eliminations, journal and journal for
Dynamics 365. Sheets are built from the pack's own files, so the workbook
can't disagree with them; nothing is recomputed. Amounts are numbers with a
number format, dates are dates, and the header row is frozen and filtered.
The CSVs don't change: they stay the files to import, and the workbook is
for reading.

The same pack gives the same workbook, byte for byte: fixed document
properties and fixed zip timestamps.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

SHEETS = ("tieout", "exceptions", "decisions", "eliminations", "journal", "journal_d365")

AMOUNTS = {
    "ledger_total",
    "trial_balance_total",
    "variance",
    "net_amount",
    "override_amount",
    "amount",
    "debit",
    "credit",
    "DEBITAMOUNT",
    "CREDITAMOUNT",
}
DATES = {"posting_date", "TRANSDATE"}
DATETIMES = {"decided_at"}
FLAGS = {"within_tolerance"}
LINE_NUMBERS = {"line", "LINENUMBER"}

FIXED_AT = datetime(1980, 1, 1)
FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)


def _typed(column: str, text: str) -> object:
    if text == "":
        return None
    if column in AMOUNTS:
        return float(Decimal(text))
    if column in LINE_NUMBERS:
        return int(text)
    if column in DATES:
        return date.fromisoformat(text)
    if column in DATETIMES:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    if column in FLAGS:
        return text == "True"
    return text


def _style(sheet: Worksheet, columns: list[str]) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        widest = max(len(str(cell.value or "")) for cell in sheet[letter][:200])
        sheet.column_dimensions[letter].width = min(max(widest + 2, 10), 60)
        if column in AMOUNTS:
            for cell in sheet[letter][1:]:
                cell.number_format = "#,##0.00"
        if column in DATES:
            for cell in sheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        if column in DATETIMES:
            for cell in sheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"


def _csv_sheet(workbook: Workbook, name: str, text: str) -> None:
    reader = csv.reader(io.StringIO(text))
    columns = next(reader)
    sheet = workbook.create_sheet(name)
    sheet.append(columns)
    for row in reader:
        sheet.append([_typed(column, value) for column, value in zip(columns, row, strict=True)])
    _style(sheet, columns)


def _run_sheet(workbook: Workbook, run: dict) -> None:
    policy, counts = run["policy"], run["counts"]
    rows: list[tuple[str, object]] = [
        ("generated_at", datetime.fromisoformat(run["generated_at"].replace("Z", "+00:00")).replace(tzinfo=None)),
        ("prepared_by", run["prepared_by"]),
        ("reviewers", ", ".join(run["reviewers"])),
        ("code_version", run["code_version"]),
        ("policy.source", policy["source"]),
        ("policy.materiality_threshold", float(Decimal(policy["materiality_threshold"]))),
        ("policy.amount_tolerance_floor", float(Decimal(policy["amount_tolerance_floor"]))),
        ("policy.amount_tolerance_rate", float(Decimal(policy["amount_tolerance_rate"]))),
        ("policy.tieout_tolerance", float(Decimal(policy["tieout_tolerance"]))),
    ]
    rows += [(f"counts.{name}", value) for name, value in counts.items()]
    rows += [(f"inputs.{entry['file']}", entry["sha256"]) for entry in run["inputs"]]
    rows += [("note", run["note"])]

    sheet = workbook.create_sheet("run")
    sheet.append(["field", "value"])
    for field, value in rows:
        sheet.append([field, value])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 80
    sheet["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value).startswith(("counts.", "policy.")) and isinstance(row[1].value, float):
            row[1].number_format = "#,##0.00########"


def _fixed_zip(data: bytes) -> bytes:
    """The same zip with every entry's timestamp fixed, so two runs of this
    function on the same workbook give the same bytes."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_DATE)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl stamps dcterms:modified at save time
                content = re.sub(
                    rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',
                    rb"\g<1>1980-01-01T00:00:00Z\g<2>",
                    content,
                )
            target.writestr(fixed, content)
    return out.getvalue()


def close_xlsx(files: dict[str, str]) -> bytes:
    """The pack's files as one workbook: the run sheet, then one per CSV."""
    run = json.loads(files["run.json"])
    workbook = Workbook()
    for blank in list(workbook.worksheets):
        workbook.remove(blank)
    _run_sheet(workbook, run)
    for name in SHEETS:
        _csv_sheet(workbook, name, files[f"{name}.csv"])
    properties = workbook.properties
    properties.creator = "intercompany-close-agent"
    properties.created = properties.modified = FIXED_AT
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _fixed_zip(buffer.getvalue())
