"""Write the close inputs as an .xlsx workbook: the template a close team fills in.

One sheet per table, named like the CSV it mirrors (entities, trial_balance,
ic_ledger), with the same columns in the same order. Amounts are numbers,
dates are dates and `interfaced` is TRUE or FALSE, so the sheets behave like
a spreadsheet rather than text. `load_dataset` reads a workbook back through
the same checks as the CSVs.
"""

from __future__ import annotations

import csv
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from intercompany_close_agent.ingest.loaders import TABLES

AMOUNTS = {"balance", "amount"}
DATES = {"transaction_date"}
FLAGS = {"interfaced"}


def _typed(column: str, text: str) -> object:
    if text == "":
        return None
    if column in AMOUNTS:
        return float(Decimal(text))
    if column in DATES:
        return date.fromisoformat(text)
    if column in FLAGS:
        return text == "True"
    return text


def write_workbook(data_dir: Path, out: Path) -> Path:
    """The folder's three CSVs as one workbook at out."""
    workbook = Workbook()
    for blank in list(workbook.worksheets):
        workbook.remove(blank)
    for table in TABLES:
        with (data_dir / f"{table}.csv").open(newline="", encoding="utf-8") as fh:
            reader = csv.DictReader(fh)
            columns = list(reader.fieldnames or [])
            sheet = workbook.create_sheet(table)
            sheet.append(columns)
            for row in reader:
                sheet.append([_typed(column, row[column]) for column in columns])

        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            letter = sheet.cell(row=1, column=index).column_letter
            widest = max(len(str(cell.value or "")) for cell in sheet[letter])
            sheet.column_dimensions[letter].width = min(max(widest + 2, 10), 60)
            if column in AMOUNTS:
                for cell in sheet[letter][1:]:
                    cell.number_format = "#,##0.00"
            if column in DATES:
                for cell in sheet[letter][1:]:
                    cell.number_format = "yyyy-mm-dd"

    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return out
