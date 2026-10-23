"""Loaders for trial balance and intercompany ledger data.

Reads the feeds a real close would pull from a consolidation tool (Lucanet,
Oracle EPM, or similar), and normalizes local GL account codes onto the
group's standard chart of accounts before anything downstream sees them.
The feeds are three CSVs in a folder, or one .xlsx workbook with the same
three tables as sheets named entities, trial_balance and ic_ledger; both go
through the same checks.

Anything wrong with a feed is raised as a DataLoadError naming the file, and
the line (or the sheet and row) where there is one. These feeds come from
outside the pipeline, and "KeyError: 'amount'" tells whoever is running the
close nothing about which of the three tables to go and fix.
"""

from __future__ import annotations

import csv
import zipfile
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from intercompany_close_agent.models.entity import Entity
from intercompany_close_agent.models.ledger import ICTransaction, TrialBalanceLine

# A handful of local-system account codes mapped onto the group's standard
# chart of accounts. Real mappings run into the thousands of lines and live
# inside the consolidation tool itself; this is a stand-in sized for the
# synthetic dataset this project ships with.
COA_MAPPING: dict[str, str] = {
    "1000-IC-REC": "IC_RECEIVABLE",
    "2000-IC-PAY": "IC_PAYABLE",
    "1100-CASH": "CASH",
    "3000-EQ": "RETAINED_EARNINGS",
    "4000-REV": "REVENUE",
    "5000-COGS": "COGS",
    "6000-OPEX": "OPEX",
}

ENTITY_COLUMNS = {"code", "name", "region", "local_currency"}
TRIAL_BALANCE_COLUMNS = {"entity_code", "account_code", "account_name", "balance"}
IC_LEDGER_COLUMNS = {"entity_code", "counterparty_code", "reference", "amount", "transaction_date"}

# The three tables, by the name of their CSV file and of their workbook sheet.
TABLES = ("entities", "trial_balance", "ic_ledger")

# (where, row): where names the file and line, or the workbook, sheet and row.
Rows = Iterator[tuple[str, dict[str, str]]]


class DataLoadError(ValueError):
    """An input file that can't be used, described well enough to go and fix it."""


def map_account_code(local_code: str) -> str:
    """Map a local GL account code onto the group standard chart of accounts.

    Falls back to the local code itself when there's no mapping entry --
    an unmapped account should surface downstream as an oddity, not vanish
    silently into the wrong bucket.
    """
    return COA_MAPPING.get(local_code, local_code)


def _rows(path: Path, required: set[str]) -> Rows:
    """(where, row) for every data row of a CSV, after checking the file is usable at all."""
    if not path.is_file():
        raise DataLoadError(f"{path} not found")
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            raise DataLoadError(f"{path.name} is empty, not even a header row")
        missing = required - set(reader.fieldnames)
        if missing:
            raise DataLoadError(f"{path.name} is missing column(s): {', '.join(sorted(missing))}")
        for row in reader:
            yield f"{path.name} line {reader.line_num}", row


AMOUNT_COLUMNS = {"amount", "balance"}


def _cell(value: object, column: str = "") -> str:
    """A workbook cell as the text the same field would have in the CSV.

    Excel keeps numbers as floats, and a whole one comes back as an int, so
    an amount arrives as 123964 or 83357.72. A float is read through its
    shortest decimal form, which is the one that was typed, and an amount
    with fewer than two places is padded to two, as the CSVs write it. Dates
    come back as datetimes and are written as the ISO date.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int | float):
        number = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
        exponent = number.as_tuple().exponent
        if column in AMOUNT_COLUMNS and isinstance(exponent, int) and exponent > -2:
            number = number.quantize(Decimal("0.01"))
        return str(number)
    return str(value)


def _sheet_rows(path: Path, sheet: str, required: set[str]) -> Rows:
    """(where, row) for every non-blank row of one sheet, with the header in row 1."""
    if not path.is_file():
        raise DataLoadError(f"{path} not found")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        raise DataLoadError(f"{path.name} isn't a readable .xlsx workbook ({exc})") from exc
    try:
        if sheet not in workbook.sheetnames:
            raise DataLoadError(f"{path.name} has no sheet named {sheet}")
        values = workbook[sheet].iter_rows(values_only=True)
        header = next(values, None)
        if header is None:
            raise DataLoadError(f"{path.name}, sheet {sheet}, is empty, not even a header row")
        names = [_cell(name) for name in header]
        missing = required - set(names)
        if missing:
            raise DataLoadError(
                f"{path.name}, sheet {sheet}, is missing column(s): {', '.join(sorted(missing))}"
            )
        for number, row in enumerate(values, start=2):
            cells = [_cell(value, column) for value, column in zip(row, names, strict=False)]
            if any(cells):
                yield f"{path.name}, sheet {sheet}, row {number}", dict(zip(names, cells, strict=False))
    finally:
        workbook.close()


def _bad_row(where: str, exc: ValueError) -> DataLoadError:
    if isinstance(exc, ValidationError):
        first = exc.errors()[0]
        field = ".".join(str(part) for part in first["loc"]) or "row"
        reason = f"{field}: {first['msg']}"
    else:
        reason = str(exc)
    return DataLoadError(f"{where}: {reason}")


def _entities(rows: Rows) -> list[Entity]:
    entities: list[Entity] = []
    for where, row in rows:
        try:
            entities.append(Entity.model_validate(row))
        except ValueError as exc:
            raise _bad_row(where, exc) from exc
    return entities


def _trial_balance(rows: Rows) -> list[TrialBalanceLine]:
    lines: list[TrialBalanceLine] = []
    for where, row in rows:
        try:
            lines.append(
                TrialBalanceLine(
                    entity_code=row["entity_code"],
                    account_code=map_account_code(row["account_code"]),
                    account_name=row["account_name"],
                    balance=_parse_decimal(row["balance"]),
                )
            )
        except ValueError as exc:
            raise _bad_row(where, exc) from exc
    return lines


def _ic_transactions(rows: Rows) -> list[ICTransaction]:
    transactions: list[ICTransaction] = []
    for where, row in rows:
        try:
            # model_validate rather than the constructor: the date arrives as text
            # and pydantic parses it, which the constructor's signature can't say.
            transactions.append(
                ICTransaction.model_validate(
                    {
                        "entity_code": row["entity_code"],
                        "counterparty_code": row["counterparty_code"],
                        "reference": row["reference"],
                        "amount": _parse_decimal(row["amount"]),
                        "currency": row.get("currency") or "USD",
                        "transaction_date": row["transaction_date"],
                        "description": row.get("description") or "",
                    }
                )
            )
        except ValueError as exc:
            raise _bad_row(where, exc) from exc
    return transactions


def load_entities(path: Path) -> list[Entity]:
    """Load the legal entity master data."""
    return _entities(_rows(path, ENTITY_COLUMNS))


def load_trial_balance(path: Path) -> list[TrialBalanceLine]:
    """Load and COA-map a trial balance CSV."""
    return _trial_balance(_rows(path, TRIAL_BALANCE_COLUMNS))


def load_ic_transactions(path: Path) -> list[ICTransaction]:
    """Load the intercompany ledger CSV into individual transaction legs."""
    return _ic_transactions(_rows(path, IC_LEDGER_COLUMNS))


def is_workbook(source: Path) -> bool:
    return source.suffix.lower() == ".xlsx"


def input_files(source: Path) -> list[Path]:
    """The files a run reads: the workbook, or the folder's three CSVs."""
    return [source] if is_workbook(source) else [source / f"{table}.csv" for table in TABLES]


def load_dataset(source: Path) -> tuple[list[Entity], list[TrialBalanceLine], list[ICTransaction]]:
    """All three feeds, checked against each other before anything runs on them.

    source is a folder with entities.csv, trial_balance.csv and ic_ledger.csv,
    or an .xlsx workbook with sheets of the same names and columns.

    Two cross-file checks, both of which would otherwise pass quietly:

    - An IC ledger with no rows. Every downstream step handles an empty list
      without complaint, so the run "succeeds" and reports a close with
      nothing to review -- exactly the output of a clean month, produced by
      a feed that never arrived.
    - An entity code the master data doesn't know. Triage reads `interfaced`
      off the master data to decide whether a small break may be
      auto-eliminated; an entity missing from it would get the benefit of
      the doubt it is least entitled to.
    """
    if is_workbook(source):
        entities = _entities(_sheet_rows(source, "entities", ENTITY_COLUMNS))
        trial_balance = _trial_balance(_sheet_rows(source, "trial_balance", TRIAL_BALANCE_COLUMNS))
        transactions = _ic_transactions(_sheet_rows(source, "ic_ledger", IC_LEDGER_COLUMNS))
        ledger, master = f"{source.name}, sheet ic_ledger,", f"{source.name}, sheet entities"
    else:
        entities = load_entities(source / "entities.csv")
        trial_balance = load_trial_balance(source / "trial_balance.csv")
        transactions = load_ic_transactions(source / "ic_ledger.csv")
        ledger, master = "ic_ledger.csv", "entities.csv"

    if not transactions:
        raise DataLoadError(f"{ledger} has no transactions, so there is nothing to close")

    known = {entity.code for entity in entities}
    used = (
        {txn.entity_code for txn in transactions}
        | {txn.counterparty_code for txn in transactions}
        | {line.entity_code for line in trial_balance}
    )
    unknown = sorted(used - known)
    if unknown:
        raise DataLoadError(f"entity code(s) not in {master}: {', '.join(unknown)}")

    return entities, trial_balance, transactions


def _parse_decimal(raw: str) -> Decimal:
    try:
        return Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"could not parse amount: {raw!r}") from exc
